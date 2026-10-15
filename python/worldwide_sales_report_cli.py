import argparse
import calendar
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


DATA_SHEET = "ข้อมูลรวม"
SUMMARY_SHEET = "สรุป"
DATA_HEADERS = [
    "รหัสรายการ", "วันที่", "บริษัท", "ร้านค้า", "จังหวัด",
    "สินค้า", "น้ำหนัก (กรัม)", "จำนวน (ลัง)", "บันทึกเมื่อ",
]
THAI_MONTHS = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม",
    "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม",
    "พฤศจิกายน", "ธันวาคม",
]
HEADER_FILL = PatternFill("solid", fgColor="6D28D9")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="E2E8F0"))


def company_label(part):
    return "Valueplus Retail" if str(part).lower() == "retail" else "Valueplus Worldwide"


def normalize_payload(payload):
    sale_id = str(payload.get("saleId") or "").strip()
    sale_date = datetime.strptime(str(payload.get("date") or ""), "%Y-%m-%d")
    store = payload.get("store") or {}
    items = payload.get("items") or []
    if not sale_id or not str(store.get("name") or "").strip() or not items:
        raise ValueError("ข้อมูลรายการขายสำหรับสร้าง Excel ไม่ครบถ้วน")

    normalized_items = []
    for item in items:
        cartons = int(item.get("cartons") or 0)
        if cartons <= 0:
            raise ValueError("จำนวนลังต้องมากกว่า 0")
        normalized_items.append({
            "name": str(item.get("name") or "").strip(),
            "weightGrams": float(item.get("weightGrams") or 0),
            "cartons": cartons,
        })
    return {
        "saleId": sale_id,
        "date": sale_date,
        "company": company_label(payload.get("part")),
        "store": str(store.get("name") or "").strip(),
        "province": str(store.get("province") or "").strip(),
        "items": normalized_items,
        "createdAt": str(payload.get("createdAt") or "").strip(),
    }


def ensure_data_sheet(workbook):
    sheet = workbook[DATA_SHEET] if DATA_SHEET in workbook.sheetnames else workbook.create_sheet(DATA_SHEET, 0)
    if sheet.max_row == 1 and not sheet.cell(1, 1).value:
        for column, value in enumerate(DATA_HEADERS, 1):
            sheet.cell(1, column, value)
    else:
        for column, value in enumerate(DATA_HEADERS, 1):
            sheet.cell(1, column, value)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(sheet.max_row, 2)}"
    widths = [24, 14, 24, 32, 18, 34, 16, 16, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    return sheet


def replace_sale_rows(sheet, record):
    for row in range(sheet.max_row, 1, -1):
        if str(sheet.cell(row, 1).value or "") == record["saleId"]:
            sheet.delete_rows(row)

    for item in record["items"]:
        sheet.append([
            record["saleId"], record["date"], record["company"],
            record["store"], record["province"], item["name"],
            item["weightGrams"], item["cartons"],
            record["createdAt"] or datetime.now(),
        ])
        row = sheet.max_row
        sheet.cell(row, 2).number_format = "dd/mm/yyyy"
        sheet.cell(row, 7).number_format = "#,##0.00"
        sheet.cell(row, 8).number_format = "#,##0"
        for cell in sheet[row]:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")
    sheet.auto_filter.ref = f"A1:I{max(sheet.max_row, 2)}"


def rebuild_summary(workbook, data_sheet, month, buddhist_year):
    if SUMMARY_SHEET in workbook.sheetnames:
        del workbook[SUMMARY_SHEET]
    sheet = workbook.create_sheet(SUMMARY_SHEET, 1)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"

    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"สรุปรายงานยอดขาย ประจำเดือน{THAI_MONTHS[month]} {buddhist_year}"
    sheet["A1"].fill = PatternFill("solid", fgColor="312E81")
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    products = sorted({str(data_sheet.cell(row, 6).value or "") for row in range(2, data_sheet.max_row + 1) if data_sheet.cell(row, 6).value})
    product_values = ["ทั้งหมด", *products]
    for row, value in enumerate(product_values, 2):
        sheet.cell(row, 10, value)
    sheet["K2"] = "ทั้งหมด"
    sheet["K3"] = "Valueplus Retail"
    sheet["K4"] = "Valueplus Worldwide"
    sheet.column_dimensions["J"].hidden = True
    sheet.column_dimensions["K"].hidden = True

    sheet["A3"] = "เลือกสินค้า"
    sheet["B3"] = "ทั้งหมด"
    sheet["A4"] = "เลือกบริษัท"
    sheet["B4"] = "ทั้งหมด"
    for address in ("A3", "A4"):
        sheet[address].font = Font(bold=True, color="475569")
    for address in ("B3", "B4"):
        sheet[address].fill = PatternFill("solid", fgColor="F0FDFA")
        sheet[address].font = Font(bold=True, color="0F766E")
        sheet[address].alignment = Alignment(horizontal="center")
    product_validation = DataValidation(
        type="list",
        formula1=f"='{SUMMARY_SHEET}'!$J$2:$J${len(product_values) + 1}",
    )
    company_validation = DataValidation(type="list", formula1=f"='{SUMMARY_SHEET}'!$K$2:$K$4")
    sheet.add_data_validation(product_validation)
    sheet.add_data_validation(company_validation)
    product_validation.add(sheet["B3"])
    company_validation.add(sheet["B4"])

    cards = [
        ("A6", "ยอดรวม (ลัง)", "B6", '=IF(AND($B$3="ทั้งหมด",$B$4="ทั้งหมด"),SUM(\'ข้อมูลรวม\'!$H:$H),IF($B$3="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$C:$C,$B$4),IF($B$4="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$F:$F,$B$3),SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$F:$F,$B$3,\'ข้อมูลรวม\'!$C:$C,$B$4))))'),
        ("C6", "จำนวนรายการ", "D6", '=IF(AND($B$3="ทั้งหมด",$B$4="ทั้งหมด"),COUNTA(\'ข้อมูลรวม\'!$A:A)-1,IF($B$3="ทั้งหมด",COUNTIF(\'ข้อมูลรวม\'!$C:C,$B$4),IF($B$4="ทั้งหมด",COUNTIF(\'ข้อมูลรวม\'!$F:F,$B$3),COUNTIFS(\'ข้อมูลรวม\'!$F:F,$B$3,\'ข้อมูลรวม\'!$C:C,$B$4))))'),
    ]
    for label_cell, label, value_cell, formula in cards:
        sheet[label_cell] = label
        sheet[value_cell] = formula
        sheet[label_cell].fill = PatternFill("solid", fgColor="EDE9FE")
        sheet[value_cell].fill = PatternFill("solid", fgColor="F5F3FF")
        sheet[label_cell].font = Font(bold=True, color="5B21B6")
        sheet[value_cell].font = Font(bold=True, color="312E81", size=15)
        sheet[value_cell].number_format = "#,##0"

    headers = ["วันที่", "ร้านค้า", "จังหวัด", "สินค้า", "จำนวน (ลัง)"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(10, column, value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    sheet["A11"] = '=IFERROR(FILTER(CHOOSE({1,2,3,4,5},\'ข้อมูลรวม\'!B2:B5000,\'ข้อมูลรวม\'!D2:D5000,\'ข้อมูลรวม\'!E2:E5000,\'ข้อมูลรวม\'!F2:F5000,\'ข้อมูลรวม\'!H2:H5000),((\'ข้อมูลรวม\'!F2:F5000=$B$3)+($B$3="ทั้งหมด"))*((\'ข้อมูลรวม\'!C2:C5000=$B$4)+($B$4="ทั้งหมด"))),"ไม่พบข้อมูล")'
    sheet["A11"].number_format = "dd/mm/yyyy"

    days_in_month = calendar.monthrange(buddhist_year - 543, month)[1]
    for day in range(1, days_in_month + 1):
        row = day + 1
        sheet.cell(row, 13, datetime(buddhist_year - 543, month, day))
        sheet.cell(row, 13).number_format = "d"
        sheet.cell(row, 14, f'=IFERROR(IF(MONTH(M{row})<>{month},0,IF(AND($B$3="ทั้งหมด",$B$4="ทั้งหมด"),SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$B:$B,M{row}),IF($B$3="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$B:$B,M{row},\'ข้อมูลรวม\'!$C:$C,$B$4),IF($B$4="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$B:$B,M{row},\'ข้อมูลรวม\'!$F:$F,$B$3),SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$B:$B,M{row},\'ข้อมูลรวม\'!$F:$F,$B$3,\'ข้อมูลรวม\'!$C:$C,$B$4))))),0)')
    sheet["M1"] = "วัน"
    sheet["N1"] = "จำนวนลัง"

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "ยอดขายรายวัน (ลัง)"
    chart.y_axis.title = "จำนวนลัง"
    chart.x_axis.title = "วันที่"
    chart.add_data(
        Reference(sheet, min_col=14, min_row=1, max_row=days_in_month + 1),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(sheet, min_col=13, min_row=2, max_row=days_in_month + 1),
    )
    chart.height = 8
    chart.width = 15
    sheet.add_chart(chart, "G3")

    sheet["P1"] = "บริษัท"
    sheet["Q1"] = "จำนวนลัง"
    sheet["P2"] = "Valueplus Retail"
    sheet["P3"] = "Valueplus Worldwide"
    sheet["Q2"] = '=IF($B$3="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$C:$C,P2),SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$C:$C,P2,\'ข้อมูลรวม\'!$F:$F,$B$3))'
    sheet["Q3"] = '=IF($B$3="ทั้งหมด",SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$C:$C,P3),SUMIFS(\'ข้อมูลรวม\'!$H:$H,\'ข้อมูลรวม\'!$C:$C,P3,\'ข้อมูลรวม\'!$F:$F,$B$3))'
    company_chart = DoughnutChart()
    company_chart.title = "สัดส่วนยอดขายตามบริษัท"
    company_chart.add_data(Reference(sheet, min_col=17, min_row=1, max_row=3), titles_from_data=True)
    company_chart.set_categories(Reference(sheet, min_col=16, min_row=2, max_row=3))
    company_chart.height = 7
    company_chart.width = 10
    sheet.add_chart(company_chart, "G19")

    for column, width in {"A": 15, "B": 28, "C": 20, "D": 32, "E": 16, "F": 3, "G": 16, "H": 16}.items():
        sheet.column_dimensions[column].width = width
    sheet.conditional_formatting.add("E11:E5000", ColorScaleRule(start_type="min", start_color="DCFCE7", end_type="max", end_color="16A34A"))
    return sheet


def update_workbook(payload, output_path):
    record = normalize_payload(payload)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(path) if path.exists() else Workbook()
    if workbook.sheetnames == ["Sheet"] and not workbook["Sheet"].cell(1, 1).value:
        del workbook["Sheet"]
    data_sheet = ensure_data_sheet(workbook)
    replace_sale_rows(data_sheet, record)
    rebuild_summary(workbook, data_sheet, record["date"].month, record["date"].year + 543)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    try:
        workbook.save(path)
    except PermissionError as error:
        raise ValueError("กรุณาปิดไฟล์ Excel รายงานยอดขายก่อนบันทึกข้อมูล") from error
    return {
        "outputPath": str(path),
        "saleId": record["saleId"],
        "rowCount": len(record["items"]),
        "totalCartons": sum(item["cartons"] for item in record["items"]),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--payload", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    try:
        payload = json.loads(args.payload)
        result = update_workbook(payload, args.output)
        print(json.dumps({"success": True, "data": result}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"success": False, "message": str(error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
