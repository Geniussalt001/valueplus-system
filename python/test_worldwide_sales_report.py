import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from worldwide_sales_report_cli import DATA_SHEET, SUMMARY_SHEET, update_workbook


class WorldwideSalesReportTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output = Path(self.temp_dir.name) / "รายงานยอดขาย" / "ตุลาคม 2569" / "รายงานยอดขาย ตุลาคม 2569.xlsx"

    def tearDown(self):
        self.temp_dir.cleanup()

    def payload(self, sale_id, cartons=3):
        return {
            "saleId": sale_id,
            "part": "retail",
            "date": "2026-10-15",
            "store": {"name": "ร้านทดสอบ", "province": "เชียงใหม่"},
            "items": [
                {"name": "สินค้า A", "weightGrams": 55, "cartons": cartons},
                {"name": "สินค้า B", "weightGrams": 70, "cartons": 2},
            ],
            "createdAt": "2026-10-15T09:00:00Z",
        }

    def test_creates_monthly_workbook_with_clean_visible_columns_and_charts(self):
        result = update_workbook(self.payload("SAL-001"), self.output)

        self.assertTrue(self.output.exists())
        self.assertEqual(result["totalCartons"], 5)
        workbook = load_workbook(self.output, data_only=False)
        self.assertEqual(workbook.sheetnames, [DATA_SHEET, SUMMARY_SHEET])
        self.assertEqual(
            [workbook[DATA_SHEET].cell(1, column).value for column in range(1, 8)],
            ["วันที่", "บริษัท", "ร้านค้า", "จังหวัด", "สินค้า", "น้ำหนัก (กรัม)", "จำนวน (ลัง)"],
        )
        self.assertTrue(workbook[DATA_SHEET].column_dimensions["H"].hidden)
        self.assertEqual(workbook[DATA_SHEET].max_row, 3)
        self.assertEqual(len(workbook[SUMMARY_SHEET]._charts), 2)
        self.assertEqual(len(workbook[SUMMARY_SHEET].data_validations.dataValidation), 0)
        self.assertEqual(workbook[SUMMARY_SHEET]["A4"].value, 5)
        self.assertFalse(any(
            cell.data_type == "f"
            for row in workbook[SUMMARY_SHEET].iter_rows()
            for cell in row
        ))

    def test_updates_same_sale_without_duplicate_rows(self):
        update_workbook(self.payload("SAL-001"), self.output)
        update_workbook(self.payload("SAL-001", cartons=8), self.output)

        workbook = load_workbook(self.output, data_only=False)
        sheet = workbook[DATA_SHEET]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["G2"].value, 8)

    def test_appends_a_different_sale(self):
        update_workbook(self.payload("SAL-001"), self.output)
        update_workbook(self.payload("SAL-002", cartons=4), self.output)

        workbook = load_workbook(self.output, data_only=False)
        self.assertEqual(workbook[DATA_SHEET].max_row, 5)

    def test_migrates_previous_report_layout(self):
        from openpyxl import Workbook

        self.output.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = DATA_SHEET
        sheet.append(["รหัสรายการ", "วันที่", "บริษัท", "ร้านค้า", "จังหวัด", "สินค้า", "น้ำหนัก (กรัม)", "จำนวน (ลัง)", "บันทึกเมื่อ"])
        sheet.append(["SAL-OLD", "15/10/2026", "Valueplus Retail", "ร้านเดิม", "กรุงเทพฯ", "สินค้าเดิม", 50, 7, "2026-10-15T09:00:00Z"])
        workbook.save(self.output)

        update_workbook(self.payload("SAL-NEW"), self.output)

        migrated = load_workbook(self.output, data_only=False)[DATA_SHEET]
        self.assertEqual(migrated["A2"].value.strftime("%d/%m/%Y"), "15/10/2026")
        self.assertEqual(migrated["E2"].value, "สินค้าเดิม")
        self.assertEqual(migrated["G2"].value, 7)
        self.assertEqual(migrated["H2"].value, "SAL-OLD")


if __name__ == "__main__":
    unittest.main()
