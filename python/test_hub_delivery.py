import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from hub_delivery_cli import handle, import_source, template_default


class HubDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.source=self.root/'Hub.xlsx';self.data=self.root/'data';self.output=self.root/'output'
        w=Workbook();s=w.active;s.title=' Hub สินค้า Dry '
        for col,sku,name in [(25,6002953,'เค้ก'),(26,6002487,'ชีส')]:s.cell(2,col,sku);s.cell(3,col,name)
        for r,code in [(4,1337),(5,2215),(6,738),(7,999)]:
            for col,value in [(1,'DRY'),(3,code),(4,'สาขา '+str(code)),(5,'WB07'),(8,1),(25,50),(26,30)]:s.cell(r,col,value)
        full=PatternFill('solid',fgColor='FF92D050')
        for col in [1,3,4,8]:s.cell(4,col).fill=full
        # A green code alone is not a ready branch.
        s.cell(5,3).fill=full
        for col in [1,3,4,8]:s.cell(6,col).fill=full
        s.cell(7,8,'เสี่ยงน้ำท่วม')
        # A later dated block repeats SKU but must not overwrite daily values.
        s.cell(2,27,6002953);s.cell(3,27,'เค้ก');s.cell(4,27,999)
        w.save(self.source)

    def tearDown(self):self.tmp.cleanup()
    def run_action(self,action,payload):return handle(action,self.data,self.output,payload)
    def test_full_row_green_and_recurring_block(self):
        rows=import_source(self.source)
        self.assertEqual([b['code'] for b in rows if b['ready']],['738','1337'])
        self.assertEqual(rows[0]['items'][0]['quantity'],50)
        self.assertTrue(next(b for b in rows if b['code']=='999')['flood'])

    def test_cached_source_export_template_and_history(self):
        self.run_action('import-source',{'path':str(self.source)})
        self.source.unlink()
        self.assertEqual(len(self.run_action('load',{})['branches']),4)
        first=self.run_action('export',{'date':'2026-10-02','codes':['738','1337']})
        output=Path(first['path']);w=load_workbook(output,data_only=True)
        self.assertEqual(w.sheetnames,['2_738','2_1337'])
        for s in w:
            self.assertEqual(s['M33'].value,80);self.assertEqual(s['N33'].value,80)
            self.assertEqual(s['M13'].value,s['N13'].value)
            self.assertEqual(s['J13'].value,s['K13'].value)
            self.assertIn('D9:P9',[str(m) for m in s.merged_cells.ranges])
            self.assertIsNone(s['N7'].value)
            self.assertEqual(s['M11'].value,'<-----ราคาขายปลีก------->')
            self.assertEqual(s.page_setup.fitToHeight,1);self.assertEqual(s.page_setup.fitToWidth,1)
        second=self.run_action('export',{'date':'2026-10-02','codes':['738']})
        self.assertNotEqual(first['path'],second['path']);self.assertTrue(output.exists())
        output.unlink();opened=self.run_action('open-history',{'id':first['store']['history'][0]['id']})
        self.assertTrue(Path(opened['path']).exists())
        self.assertEqual(len(self.run_action('load',{})['history']),2)

    def test_invalid_export_does_not_add_history(self):
        self.run_action('import-source',{'path':str(self.source)})
        for payload in [{'date':'2026-02-30','codes':['738']},{'date':'2026-10-02','codes':['999']},{'date':'2026-10-02','codes':['missing']},{'date':'2026-10-02','codes':['738','738']}]:
            with self.assertRaises(ValueError):self.run_action('export',payload)
        self.assertEqual(self.run_action('load',{})['history'],[])

    def test_settings_and_custom_template_survive_reload(self):
        self.run_action('import-source',{'path':str(self.source)})
        settings=self.run_action('load',{})['settings'];settings['preparedBy']='ผู้ทดสอบ'
        self.run_action('settings',{'settings':settings})
        self.run_action('import-template',{'path':str(template_default())})
        result=self.run_action('export',{'date':'2026-10-03','codes':['738']})
        w=load_workbook(result['path'],data_only=True)
        self.assertIn('ผู้ทดสอบ',w.worksheets[0]['B36'].value)
        self.assertIn('3 ตุลาคม 2569',w.worksheets[0]['N8'].value)
        self.assertEqual(self.run_action('load',{})['settings']['preparedBy'],'ผู้ทดสอบ')

    def test_invalid_source_keeps_previous_snapshot(self):
        original=self.run_action('import-source',{'path':str(self.source)})
        w=load_workbook(self.source);w.worksheets[0]['Y4']=-1;w.save(self.source)
        with self.assertRaises(ValueError):self.run_action('import-source',{'path':str(self.source)})
        self.assertEqual(self.run_action('load',{})['branches'],original['branches'])

if __name__=='__main__':unittest.main()
