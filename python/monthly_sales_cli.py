import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


DATE_PATTERN = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")
SALE_PREFIXES = ("IV", "IP")
MONTH_NAMES = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม",
    "มิถุนายน", "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม",
    "พฤศจิกายน", "ธันวาคม",
]


def read_csv_rows(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "cp874", "tis-620"):
        try:
            return list(csv.reader(raw.decode(encoding).splitlines())), encoding
        except UnicodeDecodeError:
            continue
    raise ValueError(f"อ่านภาษาไทยในไฟล์ CSV ไม่สำเร็จ: {path}")


def cell(row, index):
    return str(row[index] if index < len(row) else "").strip()


def number(value):
    try:
        return float(str(value or "").replace(",", "").strip())
    except ValueError:
        return 0.0


def load_products():
    path = Path(__file__).parent / "valueplus_billing" / "data" / "products.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {
        item["express_code"]: item
        for item in payload["products"]
    }


def parse_files(csv_paths):
    products = load_products()
    records = []
    excluded = defaultdict(lambda: {"quantity": 0.0, "transactions": 0, "name": ""})
    encodings = {}

    for csv_path in csv_paths:
        rows, encoding = read_csv_rows(csv_path)
        encodings[Path(csv_path).name] = encoding
        current_name = ""
        current_code = ""

        for source_row, row in enumerate(rows, 1):
            date_match = DATE_PATTERN.match(cell(row, 5))
            if not date_match:
                product_name = cell(row, 3)
                product_code = cell(row, 4)
                if product_name and product_code:
                    current_name = product_name
                    current_code = product_code
                continue

            document = cell(row, 6).upper()
            if document.startswith("SR"):
                record_type = "CN"
            elif document.startswith(SALE_PREFIXES):
                record_type = "ขาย"
            else:
                continue

            quantity = number(cell(row, 7))
            day, month, buddhist_year = map(int, date_match.groups())
            gregorian_year = buddhist_year - 543 if buddhist_year >= 2400 else buddhist_year
            product = products.get(current_code)

            if product is None:
                item = excluded[current_code or "ไม่พบรหัส"]
                item["name"] = current_name
                item["quantity"] += quantity
                item["transactions"] += 1
                continue

            records.append({
                "source_file": Path(csv_path).name,
                "source_row": source_row,
                "date": f"{day:02d}/{month:02d}/{buddhist_year:04d}",
                "year": gregorian_year,
                "buddhist_year": buddhist_year,
                "month": month,
                "month_name": MONTH_NAMES[month],
                "type": record_type,
                "express_code": current_code,
                "cpall_code": product["cpall_code"],
                "product_name": product["pdf_name"],
                "quantity": quantity,
                "document": document,
            })

    return records, dict(excluded), encodings


def build_result(csv_paths):
    records, excluded, encodings = parse_files(csv_paths)
    if not records:
        raise ValueError("ไม่พบรายการขายหรือ CN ของสินค้าที่ใช้งานในไฟล์ CSV")

    month_totals = defaultdict(lambda: {"sales": 0.0, "cn": 0.0})
    product_totals = defaultdict(lambda: {"name": "", "sales": 0.0, "cn": 0.0})
    for record in records:
        key = (record["year"], record["month"])
        target = "cn" if record["type"] == "CN" else "sales"
        month_totals[key][target] += record["quantity"]
        product = product_totals[record["cpall_code"]]
        product["name"] = record["product_name"]
        product[target] += record["quantity"]

    months = []
    for (year, month), totals in sorted(month_totals.items()):
        months.append({
            "year": year,
            "buddhist_year": year + 543,
            "month": month,
            "month_name": MONTH_NAMES[month],
            "sales": totals["sales"],
            "cn": totals["cn"],
            "net": totals["sales"] - totals["cn"],
        })

    products = []
    for code, totals in sorted(product_totals.items()):
        products.append({
            "cpall_code": code,
            "product_name": totals["name"],
            "sales": totals["sales"],
            "cn": totals["cn"],
            "net": totals["sales"] - totals["cn"],
        })

    return {
        "csv_paths": list(csv_paths),
        "file_count": len(csv_paths),
        "transaction_count": len(records),
        "product_count": len(products),
        "sales_total": sum(item["sales"] for item in months),
        "cn_total": sum(item["cn"] for item in months),
        "net_total": sum(item["net"] for item in months),
        "months": months,
        "products": products,
        "records": records,
        "excluded_products": [
            {"express_code": code, **value}
            for code, value in sorted(excluded.items())
        ],
        "excluded_count": sum(item["transactions"] for item in excluded.values()),
        "encodings": encodings,
        "output_path": "",
    }


def style_header(ws, row, start, end, color="9F1239"):
    for column in range(start, end + 1):
        cell_value = ws.cell(row, column)
        cell_value.fill = PatternFill("solid", fgColor=color)
        cell_value.font = Font(color="FFFFFF", bold=True)
        cell_value.alignment = Alignment(horizontal="center")


def export_workbook(result, output_path):
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard.merge_cells("A1:F1")
    dashboard.append(["สรุปยอดขายรายเดือน"])
    dashboard.append(["ยอดขายรวม", result["sales_total"], "CN รวม", result["cn_total"], "ยอดสุทธิ", result["net_total"]])
    dashboard.append([])
    dashboard.append(["เดือน", "ยอดขาย", "CN", "ยอดสุทธิ", "% CN"])
    for item in result["months"]:
        dashboard.append([
            f'{item["month_name"]} {item["buddhist_year"]}', item["sales"], item["cn"], item["net"],
            item["cn"] / item["sales"] if item["sales"] else 0,
        ])
    style_header(dashboard, 4, 1, 5)
    dashboard["A1"].font = Font(size=18, bold=True, color="9F1239")
    dashboard.column_dimensions["A"].width = 24
    for column in "BCDE": dashboard.column_dimensions[column].width = 16
    for row in dashboard.iter_rows(min_row=2, min_col=2, max_col=6):
        for value in row: value.number_format = "#,##0"
    for value in dashboard["E"]: value.number_format = "0.00%"
    if len(result["months"]) > 1:
        chart = LineChart()
        chart.title = "แนวโน้มยอดขายรายเดือน"
        chart.add_data(Reference(dashboard, min_col=2, max_col=4, min_row=4, max_row=4 + len(result["months"])), titles_from_data=True)
        chart.set_categories(Reference(dashboard, min_col=1, min_row=5, max_row=4 + len(result["months"])))
        chart.height = 8
        chart.width = 18
        dashboard.add_chart(chart, "G4")

    summary = wb.create_sheet("สรุปรวมรายการสินค้า")
    summary.append(["รหัสสินค้า", "ชื่อสินค้า", "ยอดขายรวม", "CN รวม", "ยอดสุทธิ", "% CN"])
    for item in result["products"]:
        summary.append([item["cpall_code"], item["product_name"], item["sales"], item["cn"], item["net"], item["cn"] / item["sales"] if item["sales"] else 0])
    style_header(summary, 1, 1, 6)
    summary.column_dimensions["A"].width = 16
    summary.column_dimensions["B"].width = 48

    data = wb.create_sheet("ฐานข้อมูลสรุป")
    data.append(["ปี", "เดือนที่", "เดือน", "ประเภท", "รหัสสินค้า", "ชื่อสินค้า", "จำนวน", "วันที่", "เลขที่เอกสาร", "ไฟล์ต้นทาง"])
    for record in result["records"]:
        data.append([record["year"], record["month"], record["month_name"], record["type"], record["cpall_code"], record["product_name"], record["quantity"], record["date"], record["document"], record["source_file"]])
    style_header(data, 1, 1, 10, "1F4E78")
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions

    if result["excluded_products"]:
        excluded = wb.create_sheet("รายการที่ไม่นำมาคำนวณ")
        excluded.append(["รหัส Express", "ชื่อสินค้า", "จำนวนธุรกรรม", "จำนวนรวม"])
        for item in result["excluded_products"]:
            excluded.append([item["express_code"], item["name"], item["transactions"], item["quantity"]])
        style_header(excluded, 1, 1, 4, "B45309")
        excluded.column_dimensions["B"].width = 52

    for ws in wb.worksheets:
        ws.freeze_panes = ws.freeze_panes or "A2"
        ws.sheet_view.showGridLines = False
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    result["output_path"] = str(output_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv", action="append", required=True)
    parser.add_argument("--output")
    parser.add_argument("--preview", action="store_true")
    args = parser.parse_args()
    try:
        result = build_result(args.csv)
        if not args.preview:
            if not args.output:
                raise ValueError("กรุณาระบุไฟล์ Excel ผลลัพธ์")
            export_workbook(result, args.output)
        response = {
            key: value
            for key, value in result.items()
            if key != "records"
        }
        print(json.dumps({"success": True, "data": response}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"success": False, "message": str(error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
