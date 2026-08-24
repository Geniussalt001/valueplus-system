import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from valueplus_summary.sales_cn_workbook import DATA_SHEET, update_sales_cn_report


class SalesCnWorkbookTest(unittest.TestCase):
    def _source(self, path: Path, rows):
        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        workbook.save(path)

    def test_adds_month_and_skips_duplicates(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "source.xlsx"
            report = folder / "report.xlsx"
            self._source(
                source,
                [
                    ["สินค้า A", "01-0000-10", 244350, "IVVPR69010001", "ลูกค้า", 100, "ชิ้น"],
                    ["สินค้า A", "01-0000-10", 244350, "SR69010001", "ลูกค้า", 5, "ชิ้น"],
                ],
            )
            first = update_sales_cn_report(str(source), str(report))
            second = update_sales_cn_report(str(source), str(report))
            self.assertEqual(first["addedRows"], 2)
            self.assertEqual(second["addedRows"], 0)
            self.assertEqual(second["skippedRows"], 2)
            self.assertEqual(first["productSummary"][0]["net"], 95)

            workbook = load_workbook(report, read_only=False, data_only=True)
            self.assertIn("มกราคม 2026", workbook.sheetnames)
            self.assertIn("ฐานสรุป", workbook.sheetnames)
            self.assertEqual(workbook[DATA_SHEET].sheet_state, "veryHidden")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
