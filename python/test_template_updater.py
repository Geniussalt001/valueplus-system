from pathlib import Path

from openpyxl import Workbook, load_workbook

from valueplus_po.template_updater import add_template_products


def _create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "สำโรง(1)"
    sheet["A8"] = 1
    sheet["B8"] = "สินค้าเดิม"
    sheet["C8"] = "=SUMIF(Data!E:E,B8,Data!F:F)"
    sheet["E8"] = "ชิ้น"
    sheet.merge_cells("A9:C9")
    sheet["A9"] = "ยอดรวมสินค้า"
    sheet["D9"] = "=SUM(D8:D8)"
    sheet["E9"] = "ชิ้น"
    sheet["B12"] = "ผู้รับผิดชอบ"
    sheet.print_area = "A1:M12"

    data = workbook.create_sheet("Data")
    data.append(["ลูกค้า", "สถานที่", "เกณฑ์", None, "สินค้า", "บรรจุ", "หน่วย"])
    data.append([None, None, None, None, "สินค้าเดิม", 12, "ชิ้น"])
    workbook.save(path)
    workbook.close()


def test_add_template_product_preserves_layout_and_creates_backup(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    _create_template(template)

    result = add_template_products(
        template,
        [{"name": "สินค้าใหม่", "product_code": "8850001", "pack_quantity": 24}],
    )

    assert Path(result["backup_path"]).is_file()
    workbook = load_workbook(template, data_only=False)
    try:
        sheet = workbook["สำโรง(1)"]
        assert sheet["B9"].value == "สินค้าใหม่"
        assert sheet["C9"].value == "=SUMIF(Data!E:E,B9,Data!F:F)"
        assert sheet["A10"].value == "ยอดรวมสินค้า"
        assert sheet["D10"].value == "=SUM(D8:D9)"
        assert sheet["B13"].value == "ผู้รับผิดชอบ"
        assert str(sheet.print_area) == "'สำโรง(1)'!$A$1:$M$13"
        assert str(next(iter(sheet.merged_cells.ranges))) == "A10:C10"

        data = workbook["Data"]
        assert data["D3"].value == "8850001"
        assert data["E3"].value == "สินค้าใหม่"
        assert data["F3"].value == 24
        assert data["G3"].value == "ชิ้น"
    finally:
        workbook.close()


def test_existing_product_gets_editable_pdf_code_without_duplicate(tmp_path: Path) -> None:
    template = tmp_path / "template.xlsx"
    _create_template(template)

    add_template_products(
        template,
        [{"name": "สินค้าเดิม", "product_code": "885ALIAS", "pack_quantity": 999}],
    )

    workbook = load_workbook(template, read_only=True, data_only=False)
    try:
        data = workbook["Data"]
        assert data.max_row == 2
        assert data["D2"].value == "885ALIAS"
        assert data["F2"].value == 12
        assert workbook["สำโรง(1)"]["A9"].value == "ยอดรวมสินค้า"
    finally:
        workbook.close()
