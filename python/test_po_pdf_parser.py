import unittest
import tempfile

from pathlib import Path

from openpyxl import Workbook

from valueplus_po.pdf_parser import (
    DATE_PATTERN,
    ITEM_PATTERN,
    _normalize_wrapped_decimal_values,
)

from valueplus_common.cpall_pdf import (
    normalize_cpall_document_date,
    repair_cpall_extracted_text,
)

from valueplus_po.template_reader import (
    read_template,
)


class PoPdfParserTest(unittest.TestCase):
    def test_template_reads_products_beyond_legacy_row_26(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "template.xlsx"
            workbook = Workbook()
            data_sheet = workbook.active
            data_sheet.title = "Data"
            data_sheet.cell(2, 5).value = (
                "ยูมิยูมิ เนโกะพุดดิ้งเค้ก 50กรัม"
            )
            target_sheet = workbook.create_sheet("มหาชัย(1)")
            target_sheet.cell(27, 2).value = (
                "ยูมิยูมิ เนโกะพุดดิ้งเค้ก 50กรัม"
            )
            target_sheet.cell(28, 2).value = "เช็คสต็อกสินค้า"
            workbook.save(path)

            catalog = read_template(path)

            self.assertEqual(
                [
                    product.row
                    for product in catalog.sheet_products["มหาชัย(1)"]
                ],
                [27],
            )

    def test_repairs_cid_marks_in_new_cpall_pdf(self):
        text = (
            "เลขที(cid:1143) : B012801967 "
            "วันที(cid:1143) : 25/08/69 "
            "คลัง BDC ส(cid:1148)าโรง คลังดี"
        )

        self.assertEqual(
            repair_cpall_extracted_text(text),
            "เลขที่ : B012801967 วันที่ : 25/08/69 "
            "คลัง BDC สำโรง คลังดี",
        )

    def test_expands_short_buddhist_year(self):
        self.assertEqual(
            normalize_cpall_document_date("25/08/69"),
            "25/08/2569",
        )

    def test_prefers_full_four_digit_buddhist_year(self):
        match = DATE_PATTERN.search(
            "วันที่ : 25/08/2569",
        )

        self.assertIsNotNone(match)
        assert match is not None
        self.assertEqual(match.group(1), "25/08/2569")
        self.assertEqual(
            normalize_cpall_document_date(match.group(1)),
            "25/08/2569",
        )

    def test_reads_seven_digit_cpall_product_code(self):
        text = (
            "1 6002424 Hมิลค์เค้กUM 55 G. "
            "1 139.00 0 10.30"
        )

        match = ITEM_PATTERN.search(text)

        self.assertIsNotNone(match)
        assert match is not None
        self.assertEqual(match.group(2), "6002424")
        self.assertEqual(match.group(4), "139.00")

    def test_joins_wrapped_thousands_quantity(self):
        text = (
            "17 8859898700316 สินค้าทดสอบ 1 1,002.    0 8.13\n"
            "                                  00\n"
        )

        normalized = _normalize_wrapped_decimal_values(text)
        match = ITEM_PATTERN.search(normalized)

        self.assertIsNotNone(match)
        assert match is not None
        self.assertEqual(match.group(4), "1,002.00")

    def test_joins_five_digit_wrapped_quantity(self):
        text = (
            "1 8859898700316 สินค้าทดสอบ 1 10,000 0 8.13\n"
            "                                  .00\n"
        )

        normalized = _normalize_wrapped_decimal_values(text)
        match = ITEM_PATTERN.search(normalized)

        self.assertIsNotNone(match)
        assert match is not None
        self.assertEqual(
            float(match.group(4).replace(",", "")),
            10000.0,
        )

    def test_does_not_join_unrelated_wrapped_decimal(self):
        text = "ยอดรวม 10,000.\n00\n"

        self.assertEqual(
            _normalize_wrapped_decimal_values(text),
            text.rstrip("\n"),
        )

    def test_keeps_complete_quantity_unchanged(self):
        text = "1 8859898700316 สินค้าทดสอบ 1 1,200.00 0 8.13\n"

        self.assertEqual(
            _normalize_wrapped_decimal_values(text),
            text.rstrip("\n"),
        )


if __name__ == "__main__":
    unittest.main()
