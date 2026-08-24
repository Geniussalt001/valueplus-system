from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


THAI_MONTHS = {
    1: "มกราคม",
    2: "กุมภาพันธ์",
    3: "มีนาคม",
    4: "เมษายน",
    5: "พฤษภาคม",
    6: "มิถุนายน",
    7: "กรกฎาคม",
    8: "สิงหาคม",
    9: "กันยายน",
    10: "ตุลาคม",
    11: "พฤศจิกายน",
    12: "ธันวาคม",
}

PRODUCTS = [
    ("6002487", "01-0000-10"),
    ("6002799", "01-0000-37"),
    ("6002425", "01-0000-38"),
    ("6002829", "01-0000-36"),
    ("6002510", "01-0000-14"),
    ("6002424", "01-0000-29"),
    ("6002798", "01-0000-34"),
    ("6002613", "01-0000-23"),
    ("6002781", "01-0000-33"),
    ("6002634", "01-0000-24"),
    ("6002797", "01-0000-35"),
    ("6002777", "01-0000-32"),
    ("6002769", "01-0000-31"),
    ("6002742", "01-0000-30"),
    ("6002567", "01-0000-16"),
    ("6002580", "01-0000-18"),
    ("6002691", "01-0000-26"),
    ("6002601", "01-0000-22"),
    ("6002690", "01-0000-27"),
]

SOURCE_TO_EXTERNAL = {source: external for external, source in PRODUCTS}
VALID_SOURCE_CODE = re.compile(r"^\d{2}-\d{4}-\d{2}$")
DATA_SHEET = "_ข้อมูลระบบ"
CHART_SHEET = "_ข้อมูลกราฟ"

RED = "C8102E"
RED_DARK = "8F0D23"
BLUE = "1F4E78"
ORANGE = "F28C28"
GREEN = "1B8A5A"
WHITE = "FFFFFF"
INK = "243247"
MUTED = "667085"
LINE = "D6DCE5"
PANEL = "F7F9FC"


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    return float(_text(value).replace(",", ""))


def _date(value: Any, document: str) -> dt.date:
    if isinstance(value, dt.datetime):
        result = value.date()
    elif isinstance(value, dt.date):
        result = value
    else:
        serial = int(float(value))
        result = dt.date(1899, 12, 30) + dt.timedelta(days=serial)

    # ไฟล์จาก Express เก็บปี พ.ศ. ใน serial/date ของ Excel
    if result.year > 2400:
        result = result.replace(year=result.year - 543)
    elif result.year > 2100:
        result = result.replace(year=result.year - 543)

    if 2000 <= result.year <= 2200:
        return result

    match = re.search(r"69(\d{2})", document)
    month = int(match.group(1)) if match else 1
    return dt.date(2026, max(1, min(month, 12)), 1)


def _fingerprint(values: Iterable[Any]) -> str:
    normalized = "\x1f".join(_text(value) for value in values)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def read_source(source_path: str) -> tuple[list[dict[str, Any]], dict[str, str]]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    worksheet = workbook.worksheets[0]

    raw_rows: list[tuple[Any, ...]] = []
    codes_by_name: dict[str, Counter[str]] = defaultdict(Counter)
    for values in worksheet.iter_rows(min_col=1, max_col=7, values_only=True):
        row = tuple(values)
        raw_rows.append(row)
        name, code = _text(row[0]), _text(row[1])
        if name and VALID_SOURCE_CODE.match(code):
            codes_by_name[name][code] += 1

    workbook.close()
    canonical = {
        name: counts.most_common(1)[0][0]
        for name, counts in codes_by_name.items()
    }

    names: dict[str, str] = {}
    transactions: list[dict[str, Any]] = []
    for row in raw_rows:
        document = _text(row[3])
        if document.startswith("IVVPR"):
            kind = "ยอดขาย"
        elif document.startswith("SR"):
            kind = "CN"
        else:
            continue

        name = _text(row[0])
        source_code = canonical.get(name, _text(row[1]))
        external_code = SOURCE_TO_EXTERNAL.get(source_code)
        if not external_code:
            continue

        try:
            qty = _number(row[5])
            transaction_date = _date(row[2], document)
        except (TypeError, ValueError, OverflowError):
            continue

        names[source_code] = name or names.get(source_code, source_code)
        identity = _fingerprint(
            [name, source_code, transaction_date.isoformat(), document, row[4], qty, row[6]]
        )
        transactions.append(
            {
                "id": identity,
                "date": transaction_date.isoformat(),
                "year": transaction_date.year,
                "month": transaction_date.month,
                "type": kind,
                "document": document,
                "customer": _text(row[4]),
                "sourceCode": source_code,
                "productCode": external_code,
                "productName": name,
                "qty": qty,
            }
        )

    return transactions, names


def _read_existing(report_path: str) -> tuple[list[dict[str, Any]], dict[str, str]]:
    if not Path(report_path).is_file():
        return [], {}

    workbook = load_workbook(report_path, read_only=True, data_only=True)
    if DATA_SHEET not in workbook.sheetnames:
        workbook.close()
        return [], {}

    worksheet = workbook[DATA_SHEET]
    transactions: list[dict[str, Any]] = []
    names: dict[str, str] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        transaction = {
            "id": _text(row[0]),
            "date": _text(row[1]),
            "year": int(row[2]),
            "month": int(row[3]),
            "type": _text(row[4]),
            "document": _text(row[5]),
            "customer": _text(row[6]),
            "sourceCode": _text(row[7]),
            "productCode": _text(row[8]),
            "productName": _text(row[9]),
            "qty": float(row[10]),
        }
        transactions.append(transaction)
        names[transaction["sourceCode"]] = transaction["productName"]
    workbook.close()
    return transactions, names


def _title(worksheet, title: str, subtitle: str = "") -> None:
    worksheet.merge_cells("A1:E2")
    cell = worksheet["A1"]
    cell.value = title
    cell.fill = PatternFill("solid", fgColor=RED_DARK)
    cell.font = Font(name="Aptos", size=22, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 16
    if subtitle:
        worksheet.merge_cells("A3:E3")
        worksheet["A3"] = subtitle
        worksheet["A3"].font = Font(name="Aptos", size=10, color=MUTED)


def _header(worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, 1):
        cell = worksheet.cell(row=row, column=column, value=label)
        cell.fill = PatternFill("solid", fgColor=RED)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _format_table(worksheet, start_row: int, end_row: int, columns: int) -> None:
    thin = Side(style="thin", color=LINE)
    for row in worksheet.iter_rows(
        min_row=start_row, max_row=end_row, min_col=1, max_col=columns
    ):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.alignment = Alignment(vertical="center")
        for cell in row[2:]:
            cell.number_format = '#,##0;[Red](#,##0);-'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=PANEL)


def _product_rows(names: dict[str, str]) -> list[tuple[str, str, str]]:
    return [
        (external, source, names.get(source, source))
        for external, source in PRODUCTS
    ]


def _aggregate(transactions: list[dict[str, Any]]):
    values: dict[tuple[int, int, str, str], float] = defaultdict(float)
    for item in transactions:
        key = (item["year"], item["month"], item["productCode"], item["type"])
        values[key] += float(item["qty"])
    return values


def _add_dashboard(
    workbook: Workbook,
    products: list[tuple[str, str, str]],
    months: list[tuple[int, int]],
    aggregate,
) -> list[dict[str, Any]]:
    worksheet = workbook.active
    worksheet.title = "Dashboard"
    _title(
        worksheet,
        "Dashboard สรุปยอดขาย / CN / ยอดสุทธิ",
        "สรุปเฉพาะ 19 รายการสินค้า — อัปเดตอัตโนมัติเมื่อเพิ่มข้อมูล",
    )
    _header(worksheet, 5, ["รหัสสินค้า", "สินค้า", "ยอดขาย", "ยอด CN", "ยอดสุทธิ"])

    summary: list[dict[str, Any]] = []
    for index, (external, _source, name) in enumerate(products, 6):
        sales = sum(aggregate.get((year, month, external, "ยอดขาย"), 0) for year, month in months)
        cn = sum(aggregate.get((year, month, external, "CN"), 0) for year, month in months)
        net = sales - cn
        worksheet.append([external, name, sales, cn, net])
        summary.append(
            {"productCode": external, "productName": name, "sales": sales, "cn": cn, "net": net}
        )

    total_row = 6 + len(products)
    worksheet.append(
        [
            "",
            "รวมทั้งหมด",
            sum(item["sales"] for item in summary),
            sum(item["cn"] for item in summary),
            sum(item["net"] for item in summary),
        ]
    )
    _format_table(worksheet, 6, total_row - 1, 5)
    for cell in worksheet[total_row]:
        cell.fill = PatternFill("solid", fgColor="FCE8EC")
        cell.font = Font(name="Aptos", bold=True, color=RED_DARK)
        cell.border = Border(top=Side(style="medium", color=RED))
    for cell in worksheet[total_row][2:]:
        cell.number_format = '#,##0;[Red](#,##0);-'

    worksheet.freeze_panes = "C6"
    worksheet.auto_filter.ref = f"A5:E{total_row}"
    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 43
    for column in "CDE":
        worksheet.column_dimensions[column].width = 17
    worksheet.sheet_view.showGridLines = False

    chart_sheet = workbook.create_sheet(CHART_SHEET)
    chart_sheet.append(["เดือน", "ยอดขาย", "ยอด CN", "ยอดสุทธิ"])
    for year, month in months:
        sales = sum(aggregate.get((year, month, external, "ยอดขาย"), 0) for external, _, _ in products)
        cn = sum(aggregate.get((year, month, external, "CN"), 0) for external, _, _ in products)
        chart_sheet.append([f"{THAI_MONTHS[month]} {year}", sales, cn, sales - cn])

    if months:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "ยอดขาย — CN — ยอดสุทธิ รายเดือน"
        chart.y_axis.title = "จำนวน (ชิ้น)"
        chart.x_axis.title = "เดือน"
        chart.height = 9
        chart.width = 19
        categories = Reference(chart_sheet, min_col=1, min_row=2, max_row=1 + len(months))
        sales = Reference(chart_sheet, min_col=2, max_col=2, min_row=1, max_row=1 + len(months))
        net = Reference(chart_sheet, min_col=4, max_col=4, min_row=1, max_row=1 + len(months))
        chart.add_data(sales, titles_from_data=True)
        chart.add_data(net, titles_from_data=True)
        chart.set_categories(categories)

        cn_chart = LineChart()
        cn = Reference(chart_sheet, min_col=3, max_col=3, min_row=1, max_row=1 + len(months))
        cn_chart.add_data(cn, titles_from_data=True)
        cn_chart.set_categories(categories)
        cn_chart.y_axis.axId = 200
        cn_chart.y_axis.title = "CN (ชิ้น)"
        cn_chart.y_axis.crosses = "max"
        cn_chart.y_axis.majorGridlines = None
        cn_chart.dataLabels = DataLabelList()
        cn_chart.dataLabels.showVal = False
        chart += cn_chart
        worksheet.add_chart(chart, "G5")

    chart_sheet.sheet_state = "hidden"
    return summary


def _add_month_sheets(workbook: Workbook, products, months, aggregate) -> None:
    for year, month in months:
        worksheet = workbook.create_sheet(f"{THAI_MONTHS[month]} {year}")
        _title(
            worksheet,
            f"สรุปยอดประจำเดือน{THAI_MONTHS[month]} {year}",
            "ยอดสุทธิ = ยอดขาย - ยอด CN",
        )
        _header(worksheet, 5, ["รหัสสินค้า", "สินค้า", "ยอดขาย", "ยอด CN", "ยอดสุทธิ"])
        for external, _source, name in products:
            sales = aggregate.get((year, month, external, "ยอดขาย"), 0)
            cn = aggregate.get((year, month, external, "CN"), 0)
            worksheet.append([external, name, sales, cn, sales - cn])
        total_row = 6 + len(products)
        worksheet.append(
            [
                "",
                "รวมเดือน",
                sum(worksheet.cell(row=row, column=3).value for row in range(6, total_row)),
                sum(worksheet.cell(row=row, column=4).value for row in range(6, total_row)),
                sum(worksheet.cell(row=row, column=5).value for row in range(6, total_row)),
            ]
        )
        _format_table(worksheet, 6, total_row - 1, 5)
        for cell in worksheet[total_row]:
            cell.fill = PatternFill("solid", fgColor="FCE8EC")
            cell.font = Font(name="Aptos", bold=True, color=RED_DARK)
        for cell in worksheet[total_row][2:]:
            cell.number_format = '#,##0;[Red](#,##0);-'
        worksheet.freeze_panes = "C6"
        worksheet.auto_filter.ref = f"A5:E{total_row}"
        worksheet.column_dimensions["A"].width = 16
        worksheet.column_dimensions["B"].width = 43
        for column in "CDE":
            worksheet.column_dimensions[column].width = 17
        worksheet.sheet_view.showGridLines = False


def _add_base_sheet(workbook: Workbook, products, months, aggregate) -> None:
    worksheet = workbook.create_sheet("ฐานสรุป")
    _title(worksheet, "ฐานสรุปยอดรายเดือน", "ใช้ตรวจสอบยอดของแต่ละสินค้าในทุกเดือน")
    _header(worksheet, 5, ["เดือน", "รหัสสินค้า", "สินค้า", "ยอดขาย", "ยอด CN", "ยอดสุทธิ"])
    row = 6
    for year, month in months:
        for external, _source, name in products:
            sales = aggregate.get((year, month, external, "ยอดขาย"), 0)
            cn = aggregate.get((year, month, external, "CN"), 0)
            worksheet.cell(row=row, column=1, value=f"{THAI_MONTHS[month]} {year}")
            worksheet.cell(row=row, column=2, value=external)
            worksheet.cell(row=row, column=3, value=name)
            worksheet.cell(row=row, column=4, value=sales)
            worksheet.cell(row=row, column=5, value=cn)
            worksheet.cell(row=row, column=6, value=sales - cn)
            row += 1
    if row > 6:
        _format_table(worksheet, 6, row - 1, 6)
        worksheet.auto_filter.ref = f"A5:F{row - 1}"
    worksheet.freeze_panes = "D6"
    widths = [20, 16, 43, 17, 17, 17]
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.sheet_view.showGridLines = False


def _add_data_sheet(workbook: Workbook, transactions: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(DATA_SHEET)
    worksheet.append(
        [
            "id",
            "date",
            "year",
            "month",
            "type",
            "document",
            "customer",
            "sourceCode",
            "productCode",
            "productName",
            "qty",
        ]
    )
    for item in transactions:
        worksheet.append(
            [
                item["id"],
                item["date"],
                item["year"],
                item["month"],
                item["type"],
                item["document"],
                item["customer"],
                item["sourceCode"],
                item["productCode"],
                item["productName"],
                item["qty"],
            ]
        )
    worksheet.sheet_state = "veryHidden"


def update_sales_cn_report(source_path: str, report_path: str) -> dict[str, Any]:
    new_transactions, new_names = read_source(source_path)
    existing_transactions, existing_names = _read_existing(report_path)

    existing_ids = {item["id"] for item in existing_transactions}
    added = [item for item in new_transactions if item["id"] not in existing_ids]
    skipped = len(new_transactions) - len(added)
    transactions = existing_transactions + added
    if not transactions:
        raise ValueError("ไม่พบรายการขายหรือ CN ของสินค้า 19 รายการในไฟล์ที่เลือก")

    names = {**existing_names, **new_names}
    for item in transactions:
        if item["productName"]:
            names[item["sourceCode"]] = item["productName"]

    months = sorted({(int(item["year"]), int(item["month"])) for item in transactions})
    products = _product_rows(names)
    aggregate = _aggregate(transactions)

    workbook = Workbook()
    summary = _add_dashboard(workbook, products, months, aggregate)
    _add_month_sheets(workbook, products, months, aggregate)
    _add_base_sheet(workbook, products, months, aggregate)
    _add_data_sheet(workbook, transactions)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    destination = Path(report_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp_handle, temp_name = tempfile.mkstemp(
        prefix="valueplus-sales-cn-", suffix=".xlsx", dir=str(destination.parent)
    )
    os.close(temp_handle)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, destination)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)

    return {
        "reportPath": str(destination),
        "addedRows": len(added),
        "skippedRows": skipped,
        "totalRows": len(transactions),
        "months": [f"{THAI_MONTHS[month]} {year}" for year, month in months],
        "productSummary": summary,
        "message": (
            f"เพิ่มข้อมูลใหม่ {len(added):,} รายการ "
            f"ข้ามข้อมูลซ้ำ {skipped:,} รายการ และอัปเดตรายงานเรียบร้อย"
        ),
    }
