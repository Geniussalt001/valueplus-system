import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_HEADERS = {
    "รหัสสาขา", "ชื่อสาขา", "จังหวัด", "ภาค",
}


def text(value):
    return str(value or "").strip()


def branch_code(value):
    raw = text(value)
    if raw.endswith(".0"):
        raw = raw[:-2]
    digits = "".join(character for character in raw if character.isdigit())
    return digits.zfill(5) if digits else ""


def build_result(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "ข้อมูลทั้งหมด" not in workbook.sheetnames:
        raise ValueError("ไม่พบชีต 'ข้อมูลทั้งหมด' ในไฟล์ Master")
    sheet = workbook["ข้อมูลทั้งหมด"]
    rows = sheet.iter_rows(values_only=True)
    headers = [text(value) for value in next(rows, [])]
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise ValueError("หัวตาราง Master ไม่ครบ: " + ", ".join(sorted(missing)))
    positions = {header: headers.index(header) for header in REQUIRED_HEADERS}
    unique = {}
    duplicate_count = 0
    conflict_count = 0
    for source_row, row in enumerate(rows, 2):
        code = branch_code(row[positions["รหัสสาขา"]])
        if not code:
            continue
        record = {
            "branch_code": code,
            "branch_name": text(row[positions["ชื่อสาขา"]]),
            "province": text(row[positions["จังหวัด"]]),
            "region": text(row[positions["ภาค"]]),
            "source_row": source_row,
        }
        if not all(record[key] for key in ("branch_name", "province", "region")):
            raise ValueError(f"ข้อมูลสาขาไม่ครบที่แถว {source_row}")
        if code in unique:
            duplicate_count += 1
            previous = unique[code]
            if any(previous[key] != record[key] for key in ("branch_name", "province", "region")):
                conflict_count += 1
            continue
        unique[code] = record
    if conflict_count:
        raise ValueError(f"พบรหัสสาขาซ้ำแต่ข้อมูลขัดแย้ง {conflict_count} รหัส")
    records = [unique[code] for code in sorted(unique)]
    regions = {}
    provinces = set()
    for record in records:
        regions[record["region"]] = regions.get(record["region"], 0) + 1
        provinces.add(record["province"])
    return {
        "file_name": Path(path).name,
        "branch_count": len(records),
        "province_count": len(provinces),
        "duplicate_count": duplicate_count,
        "conflict_count": conflict_count,
        "regions": [{"region": key, "branch_count": regions[key]} for key in sorted(regions)],
        "records": records,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    args = parser.parse_args()
    try:
        print(json.dumps({"success": True, "data": build_result(args.xlsx)}, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({"success": False, "message": str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
